from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated, Any
import io
import re

from fastapi import APIRouter, HTTPException
from fastapi.responses import HTMLResponse, PlainTextResponse, StreamingResponse
from pydantic import BaseModel, ConfigDict, Field, StringConstraints, field_validator

from app.analysis.quality import assess_quality
from app.analysis.helpers import analyze_if_missing
from app.api.deps import get_store

router = APIRouter(prefix="/api/reports", tags=["reports"])
SpectrumId = Annotated[str, StringConstraints(pattern=r"^[A-Za-z0-9_-]{1,64}$")]


class ReportRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    ids: list[SpectrumId] = Field(min_length=1, max_length=100)
    title: Annotated[str, StringConstraints(strip_whitespace=True, max_length=200)] = (
        "ChemApp Analysis Report"
    )
    include_peaks: bool = True

    @field_validator("ids")
    @classmethod
    def unique_ids(cls, value):
        if len(value) != len(set(value)):
            raise ValueError("Spectrum IDs must be unique")
        return value

    @field_validator("title")
    @classmethod
    def single_line_title(cls, value):
        if any(ord(character) < 32 for character in value):
            raise ValueError("Report title must be a single line")
        return value or "ChemApp Analysis Report"


def _ensure_result(stored):
    if stored.result is not None:
        return stored.result

    result = analyze_if_missing(stored.spectrum)
    if result is None:
        return None
    result.metrics["quality"] = assess_quality(stored.spectrum, result)
    # Report generation is read-only: do not create or overwrite a persisted
    # analysis merely because the user exported a report.
    return result


def _format_metric_value(value: Any) -> str:
    if isinstance(value, float):
        return f"{value:.4g}"
    if isinstance(value, (list, tuple)):
        return ", ".join(_format_metric_value(v) for v in value[:8])
    if isinstance(value, dict):
        return ", ".join(f"{k}={_format_metric_value(v)}" for k, v in list(value.items())[:6])
    return str(value)


@router.post("/markdown")
def export_markdown_report(payload: ReportRequest):
    markdown = _build_markdown(payload.ids, payload.title, payload.include_peaks)
    filename = "chemapp-report.md"
    return PlainTextResponse(
        markdown,
        media_type="text/markdown",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _build_markdown(ids: list[str], title: str, include_peaks: bool = True) -> str:
    if not ids:
        raise HTTPException(400, "Provide at least one spectrum ID in 'ids'")
    store = get_store()
    lines = [
        f"# {title}",
        "",
        f"Generated: {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
        "",
    ]

    for sid in ids:
        stored = store.get(sid)
        if stored is None:
            raise HTTPException(404, f"Spectrum {sid} not found")

        spectrum = stored.spectrum
        result = _ensure_result(stored)
        name = Path(spectrum.source_file).name if spectrum.source_file else sid
        lines.extend([
            f"## {spectrum.technique.value}: {name}",
            "",
            f"- ID: `{sid}`",
            f"- Points: {spectrum.num_points}",
            f"- X range: {spectrum.x_range[0]:.4g} to {spectrum.x_range[1]:.4g} {spectrum.x_unit}",
        ])

        if result is None:
            lines.extend(["", "No analysis result is available.", ""])
            continue

        quality = result.metrics.get("quality") or assess_quality(spectrum, result)
        review_lines = [
            f"- Quality: {quality.get('status')} ({float(quality.get('score', 0)) * 100:.0f}%)",
            f"- Review status: {'Manual confirmed' if result.metrics.get('manual_confirmed') else 'Automatic analysis'}",
        ]
        if result.metrics.get("manual_confirmed"):
            review_lines.append(f"- Manual version: v{result.metrics.get('manual_version', 1)}")
        if result.metrics.get("restored_from_version"):
            review_lines.append(f"- Restored from version: v{result.metrics.get('restored_from_version')}")

        lines.extend([
            *review_lines,
            "",
            "### Summary",
            "",
            result.summary or "No summary available.",
            "",
            "### Metrics",
            "",
        ])

        for key, value in result.metrics.items():
            if key in {"quality", "channel_peaks", "segments_meta"}:
                continue
            lines.append(f"- {key}: {_format_metric_value(value)}")

        warnings = quality.get("warnings") or []
        info = quality.get("info") or []
        if warnings or info:
            lines.extend(["", "### Data Quality", ""])
            for warning in warnings:
                lines.append(f"- Warning: {warning}")
            for item in info:
                lines.append(f"- Note: {item}")

        if include_peaks and result.peaks:
            lines.extend(["", "### Top Peaks", "", "| Position | Intensity | Area | Width |", "|---:|---:|---:|---:|"])
            peaks = sorted(result.peaks, key=lambda p: abs(p.intensity), reverse=True)[:25]
            for peak in peaks:
                area = "" if peak.area is None else f"{peak.area:.4g}"
                width = "" if peak.width is None else f"{peak.width:.4g}"
                lines.append(f"| {peak.position:.4g} | {peak.intensity:.4g} | {area} | {width} |")

        if spectrum.technique.value == "XRD":
            _append_xrd_sections(lines, result)
        if spectrum.technique.value == "HPLC":
            _append_hplc_sections(lines, result)

        lines.append("")

    return "\n".join(lines)


@router.post("/html")
def export_html_report(payload: ReportRequest):
    markdown = _build_markdown(payload.ids, payload.title, payload.include_peaks)
    html = _markdown_to_simple_html(markdown, payload.title)
    return HTMLResponse(
        html,
        headers={"Content-Disposition": "attachment; filename=chemapp-report.html"},
    )


@router.post("/docx")
def export_docx_report(payload: ReportRequest):
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except Exception as exc:
        raise HTTPException(500, f"python-docx is unavailable: {exc}")

    markdown = _build_markdown(payload.ids, payload.title, payload.include_peaks)

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    current_table: list[list[str]] = []

    def flush_table():
        nonlocal current_table
        if not current_table:
            return
        rows = [row for row in current_table if not all(re.fullmatch(r"-+", c.strip()) for c in row)]
        current_table = []
        if not rows:
            return
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, cell in enumerate(hdr):
            cell.text = rows[0][i]
        for row in rows[1:]:
            cells = table.add_row().cells
            for i, value in enumerate(row[:len(cells)]):
                cells[i].text = value

    for line in markdown.splitlines():
        if line.startswith("|") and line.endswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            current_table.append(cells)
            continue
        flush_table()
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line.strip())
    flush_table()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=chemapp-report.docx"},
    )


def _markdown_to_simple_html(markdown: str, title: str) -> str:
    body = []
    in_table = False
    for line in markdown.splitlines():
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                continue
            if not in_table:
                body.append("<table>")
                in_table = True
                tag = "th"
            else:
                tag = "td"
            body.append("<tr>" + "".join(f"<{tag}>{_escape_html(c)}</{tag}>" for c in cells) + "</tr>")
            continue
        if in_table:
            body.append("</table>")
            in_table = False
        if line.startswith("# "):
            body.append(f"<h1>{_escape_html(line[2:].strip())}</h1>")
        elif line.startswith("## "):
            body.append(f"<h2>{_escape_html(line[3:].strip())}</h2>")
        elif line.startswith("### "):
            body.append(f"<h3>{_escape_html(line[4:].strip())}</h3>")
        elif line.startswith("#### "):
            body.append(f"<h4>{_escape_html(line[5:].strip())}</h4>")
        elif line.startswith("- "):
            body.append(f"<p class='bullet'>• {_escape_html(line[2:].strip())}</p>")
        elif line.strip():
            body.append(f"<p>{_escape_html(line.strip())}</p>")
    if in_table:
        body.append("</table>")
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{_escape_html(title)}</title>
<style>
body {{ font-family: Arial, sans-serif; margin: 32px; color: #18212f; }}
h1 {{ font-size: 26px; margin-bottom: 8px; }}
h2 {{ border-top: 1px solid #d7deea; padding-top: 18px; margin-top: 26px; }}
h3 {{ margin-top: 20px; color: #334155; }}
p {{ line-height: 1.45; }}
.bullet {{ margin-left: 14px; }}
table {{ width: 100%; border-collapse: collapse; margin: 10px 0 18px; font-size: 12px; }}
th, td {{ border: 1px solid #d7deea; padding: 6px 8px; text-align: left; }}
th {{ background: #eef3fb; }}
</style></head><body>{''.join(body)}</body></html>"""


def _escape_html(value: Any) -> str:
    return (
        str(value)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _append_xrd_sections(lines: list[str], result) -> None:
    phase_matches = getattr(result, "phase_matches", []) or []
    if phase_matches:
        lines.extend([
            "",
            "### XRD Qualitative Phase Analysis",
            "",
            "| Phase | Formula | Card | System | Matched | Score |",
            "|---|---|---|---|---:|---:|",
        ])
        for row in phase_matches[:8]:
            lines.append(
                f"| {row['phase_name']} | {row['formula']} | {row['card_number']} | "
                f"{row['crystal_system']} | {row['matched_peaks']}/{row['reference_peaks']} | {row['match_score']:.1f}% |"
            )

    assignments = getattr(result, "peak_assignments", []) or []
    if assignments:
        lines.extend([
            "",
            "### XRD Peak Shift / Assignment",
            "",
            "| Obs. 2θ | Ref. 2θ | Shift | hkl | Phase | Rel. I |",
            "|---:|---:|---:|---|---|---:|",
        ])
        for row in assignments[:25]:
            lines.append(
                f"| {row['observed_two_theta']:.3f} | {row['reference_two_theta']:.3f} | "
                f"{row['peak_shift']:.4f} | {row['hkl']} | {row['phase_name']} | {row['relative_intensity']:.1f}% |"
            )

    lattice = getattr(result, "lattice_parameters", []) or []
    if lattice:
        lines.extend([
            "",
            "### XRD Lattice Parameters",
            "",
            "| Phase | System | a (Å) | c (Å) | c/a | Indexed |",
            "|---|---|---:|---:|---:|---:|",
        ])
        for row in lattice:
            lines.append(
                f"| {row['phase_name']} | {row['crystal_system']} | "
                f"{_format_metric_value(row.get('a_angstrom'))} | {_format_metric_value(row.get('c_angstrom'))} | "
                f"{_format_metric_value(row.get('c_over_a'))} | {row['indexed_peaks']} |"
            )

    crystallinity = getattr(result, "crystallinity", {}) or {}
    williamson_hall = getattr(result, "williamson_hall", {}) or {}
    size_distribution = getattr(result, "size_distribution", {}) or {}
    if crystallinity or williamson_hall or size_distribution:
        lines.extend(["", "### XRD Crystallinity / Size", ""])
        if crystallinity:
            lines.append(f"- Crystallinity: {crystallinity['crystallinity_percent']:.1f}% ({crystallinity['method']})")
        if williamson_hall:
            lines.append(
                f"- Williamson-Hall: size {williamson_hall['crystallite_size_a']:.1f} Å, "
                f"microstrain {williamson_hall['microstrain']:.4g}, R² {williamson_hall['r_squared']:.3f}"
            )
        if size_distribution:
            lines.append(
                f"- Scherrer distribution: mean {size_distribution['mean_a']:.1f} Å, "
                f"median {size_distribution['median_a']:.1f} Å, std {size_distribution['std_a']:.1f} Å"
            )

    quantitative = getattr(result, "quantitative_analysis", []) or []
    if quantitative:
        lines.extend([
            "",
            "### XRD Quantitative Analysis (RIR)",
            "",
            "| Phase | Formula | RIR | Weight % |",
            "|---|---|---:|---:|",
        ])
        for row in quantitative:
            lines.append(f"| {row['phase_name']} | {row['formula']} | {row['rir']:.2f} | {row['weight_percent']:.1f}% |")


def _append_hplc_sections(lines: list[str], result) -> None:
    channel_peaks = result.metrics.get("channel_peaks") or {}
    if not channel_peaks:
        return

    lines.extend(["", "### HPLC Peak Tables", ""])
    for channel_name, channel_data in channel_peaks.items():
        wavelength = channel_data.get("wavelength_nm")
        total_area = channel_data.get("total_area", 0)
        label = f"{channel_name} ({wavelength:g} nm)" if isinstance(wavelength, (int, float)) else channel_name
        lines.extend([
            f"#### {label}",
            "",
            f"- Total area: {_format_metric_value(total_area)}",
            f"- Source: {channel_data.get('source', 'unknown')}",
            "",
            "| Retention time (min) | Type | Width (min) | Area | Height (mAU) | Area % | Name |",
            "|---:|---|---:|---:|---:|---:|---|",
        ])
        for peak in channel_data.get("peaks", []):
            lines.append(
                f"| {peak.get('position', 0):.3f} | {peak.get('type', '')} | "
                f"{peak.get('width', 0):.2f} | {peak.get('area', 0):.2f} | "
                f"{peak.get('intensity', 0):.2f} | {peak.get('area_percent', 0):.2f} | "
                f"{peak.get('name', '')} |"
            )
        lines.append("")
